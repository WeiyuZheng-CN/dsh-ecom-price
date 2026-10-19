"""ladder.py — 把 audit.mjs 抓到的 SKU 价整理成「每个规格一个价」的对照表。

    python scripts/ladder.py --aud aud.json --xlsx 采购清单.xlsx [--out ladder.md]

做三件事：
  1. 用清单里的超链接把行对到 aud.json 里那条商品（按 taobao/tmall 的 id= 或京东的数字 sku）
  2. 把该链接的 SKU 价按规格归类（ml / L / kg / g / 寸 / cm / 支 / 卷 / 只 / 个 / 张 / 片 / 块 / 把 / 套）
  3. 标出「数量 × 单价 ≠ 小计」这种表内自相矛盾

它只做归类与对照，**不替你判定对错**——判定靠人看「表里的单价落在哪个规格的价档上」。
"""
from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict

from openpyxl import load_workbook

SIZE_UNITS = [
    "毫升", "千克", "英寸",
    "ml", "mL", "ML", "kg", "KG", "mm", "cm", "L", "升", "g", "G", "克", "米", "m", "寸", "号",
    "支", "卷", "只", "个", "张", "片", "块", "把", "套", "双", "条",
]
# 长的单位必须排在短的前面（mm 在 m 前、ml 在 m 前），否则 "24mm" 会被切成 "24m"
SIZE_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(" + "|".join(re.escape(u) for u in SIZE_UNITS) + r")")
PACK_RE = re.compile(r"(\d+)\s*(件装|个装|只装|卷装|片装|支装)")
# 多色套装 / 工具箱：这类 SKU 的价格是「一整套」的，不能当单件价比较
SET_RE = re.compile(r"\d+\s*色|件套|工具箱|材料包|工具全套")

UNIT_ORDER = SIZE_UNITS + ["件装", "个装", "只装", "卷装", "片装", "支装"]


def id_from_link(link: str) -> str | None:
    if not link:
        return None
    m = re.search(r"[?&]id=(\d+)", link)
    if m:
        return m.group(1)
    m = re.search(r"item\.jd\.com/(\d+)\.html", link)
    if m:
        return m.group(1)
    return None


def size_tokens(text: str) -> list[str]:
    """Extract comparable size tokens: '500ml', '3寸', '10L', '50个装' …"""
    text = (text or "").replace("1/2寸", "0.5寸")
    text = re.sub(r"(\d+)\s*#", r"\1号", text)  # 笔号常写成 "0#"、"2#"
    out = [m.group(1) + m.group(2) for m in SIZE_RE.finditer(text)]
    out += [m.group(1) + m.group(2) for m in PACK_RE.finditer(text)]
    # 去重但保序
    seen, uniq = set(), []
    for t in out:
        if t not in seen:
            seen.add(t)
            uniq.append(t)
    return uniq


def sort_key(label: str):
    """同单位按数值排，不同单位按 UNIT_ORDER 排。"""
    tokens = size_tokens(label)
    tok = tokens[0] if tokens else ""
    m = re.match(r"(\d+(?:\.\d+)?)(.*)", tok)
    num, unit = (float(m.group(1)), m.group(2)) if m else (0.0, tok)
    rank = UNIT_ORDER.index(unit) if unit in UNIT_ORDER else len(UNIT_ORDER)
    return (rank, num, label)


def norm_price(p) -> float | None:
    try:
        return float(str(p).replace(",", ""))
    except (TypeError, ValueError):
        return None


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--aud", required=True, help="audit.mjs 输出的 json")
    ap.add_argument("--xlsx", required=True, help="采购清单 xlsx")
    ap.add_argument("--sheet", default=None, help="只处理某个工作表")
    ap.add_argument("--map", default=None, help='短链/自定义 key 时手动对应：{"清单物品名": "aud.json 里的 key"}')
    ap.add_argument("--out", default=None, help="写成 markdown 文件")
    args = ap.parse_args()

    aud = json.load(open(args.aud, encoding="utf-8"))
    manual = json.load(open(args.map, encoding="utf-8")) if args.map else {}
    by_id = {}
    for key, rec in aud.items():
        if not isinstance(rec, dict):
            continue
        for cand in (rec.get("itemId"), id_from_link(rec.get("url", "")), id_from_link(rec.get("srcUrl", ""))):
            if cand:
                by_id.setdefault(str(cand), (key, rec))

    def resolve(name: str, link: str | None):
        if name in manual and manual[name] in aud:
            return manual[name], aud[manual[name]]
        iid = id_from_link(link or "")
        if iid and iid in by_id:
            return by_id[iid]
        base = (link or "").split("?")[0]
        if base:
            for k, rec in aud.items():
                if isinstance(rec, dict) and (base in str(rec.get("url", "")) or base in str(rec.get("srcUrl", ""))):
                    return k, rec
        return None

    wb = load_workbook(args.xlsx)
    lines: list[str] = []
    sheets = [args.sheet] if args.sheet else wb.sheetnames
    problems = 0

    for sn in sheets:
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        lines.append(f"## {sn}")
        for r in range(2, ws.max_row + 1):
            name = ws.cell(r, 3).value
            if not name:
                continue
            spec = str(ws.cell(r, 4).value or "")
            qty, unit, sub = ws.cell(r, 5).value, ws.cell(r, 6).value, ws.cell(r, 7).value
            link = ws.cell(r, 9).hyperlink.target if ws.cell(r, 9).hyperlink else None
            iid = id_from_link(link or "")
            flag = ""
            if isinstance(qty, (int, float)) and isinstance(unit, (int, float)):
                calc = round(qty * unit, 2)
                if isinstance(sub, (int, float)) and abs(calc - sub) > 0.011:
                    flag = f" ⚠️ 数量×单价={calc} 但小计={sub}"
                    problems += 1
            lines.append(f"\n### {name}｜{spec}｜数量 {qty}｜单价 {unit}{flag}")

            hit = resolve(str(name).strip(), link)
            if not hit:
                lines.append("- （这个链接没有 SKU 数据：未抓 / 抓取失败 / 短链没对上；短链或自定义 key 用 --map 指定）")
                continue
            key, rec = hit
            if not rec.get("ok"):
                lines.append(f"- 抓取失败：{rec.get('reason')} {rec.get('hint') or ''}".rstrip())
                continue

            # 规格 -> Counter(价格)
            want = size_tokens(spec)
            buckets: dict[str, Counter] = defaultdict(Counter)
            for sku in rec.get("skus") or []:
                price = sku.get("priceText") or sku.get("price")
                if not price:
                    continue
                attrs = sku.get("attrsText") or sku.get("attrs") or ""
                toks = size_tokens(attrs)
                # 多色套装/工具箱的价格不能当单件价看，单独放一栏
                if SET_RE.search(attrs) and not any(t in attrs for t in want):
                    buckets["套装/多色"][str(price)] += 1
                    continue
                label = next((t for t in toks if t in want), toks[0] if toks else "（无规格）")
                buckets[label][str(price)] += 1

            lines.append(
                f"- 页面 {rec.get('skuCount')} 个 SKU，打码 {rec.get('maskedCount') or 0} 个"
                f"，可读价区间 {rec.get('minReal')}–{rec.get('maxReal')}"
            )
            if want:
                lines.append(f"- 清单规格里的量词：{'、'.join(want)}")
            for label in sorted(buckets, key=sort_key):
                counter = buckets[label]
                total = sum(counter.values())
                parts = [f"{p}（{n} 个）" if total > 1 else p for p, n in counter.most_common(4)]
                more = "…" if len(counter) > 4 else ""
                mark = " ← 与清单规格同名" if any(w in label for w in want) else ""
                lines.append(f"  - {label}: {'、'.join(parts)}{more}{mark}")

    text = "\n".join(lines)
    print(text if len(text) < 6000 else text[:6000] + "\n…（已截断，用 --out 看全文）")
    if args.out:
        open(args.out, "w", encoding="utf-8").write(text)
        print(f"\n[written] {args.out}  表内矛盾 {problems} 处")


if __name__ == "__main__":
    main()
