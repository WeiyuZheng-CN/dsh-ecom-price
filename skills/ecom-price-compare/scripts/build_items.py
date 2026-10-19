#!/usr/bin/env python
"""Turn a procurement spreadsheet into the items.json the scrapers eat.

Default column layout matches the common Chinese 采购清单 sheet:
  A 序号 | B 类别 | C 物品 | D 规格 | E 数量 | F 参考单价 | I 淘宝链接 | J 京东链接
Override with --col-* if the sheet differs. Hyperlinks on the link columns are
read for their q= / keyword= parameter, which is reused as the search query.

  python build_items.py --xlsx 清单.xlsx --out items.json
  python build_items.py --xlsx 清单.xlsx --sheet Sheet1 --start-row 3 --out items.json
"""
import argparse
import json
import urllib.parse

from openpyxl import load_workbook


def col_index(letter):
    n = 0
    for ch in letter.strip().upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def query_of(hyperlink, key):
    if not hyperlink or not hyperlink.target:
        return None
    q = urllib.parse.urlparse(hyperlink.target).query
    return urllib.parse.parse_qs(q).get(key, [None])[0]


def numeric(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace(",", "").strip())
        except ValueError:
            return None
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sheet", default=None, help="default: first sheet")
    ap.add_argument("--start-row", type=int, default=3)
    ap.add_argument("--end-row", type=int, default=0, help="0 = to the end")
    ap.add_argument("--out", default="items.json")
    ap.add_argument("--col-id", default="A")
    ap.add_argument("--col-cat", default="B")
    ap.add_argument("--col-name", default="C")
    ap.add_argument("--col-spec", default="D")
    ap.add_argument("--col-qty", default="E")
    ap.add_argument("--col-ref", default="F")
    ap.add_argument("--col-tb", default="I")
    ap.add_argument("--col-jd", default="J")
    args = ap.parse_args()

    wb = load_workbook(args.xlsx, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb.worksheets[0]

    c_id, c_cat, c_name = col_index(args.col_id), col_index(args.col_cat), col_index(args.col_name)
    c_spec, c_qty, c_ref = col_index(args.col_spec), col_index(args.col_qty), col_index(args.col_ref)
    c_tb, c_jd = col_index(args.col_tb), col_index(args.col_jd)

    end = args.end_row or ws.max_row
    items = []
    for row in range(args.start_row, end + 1):
        idx = ws.cell(row, c_id).value
        if idx is None:
            continue
        try:
            idx = int(idx)
        except (TypeError, ValueError):
            continue
        tb = ws.cell(row, c_tb).hyperlink
        jd = ws.cell(row, c_jd).hyperlink
        items.append(
            {
                "id": idx,
                "cat": ws.cell(row, c_cat).value,
                "name": ws.cell(row, c_name).value,
                "spec": ws.cell(row, c_spec).value,
                "qty": ws.cell(row, c_qty).value,
                "refPrice": numeric(ws.cell(row, c_ref).value),
                "tbQuery": query_of(tb, "q"),
                "jdQuery": query_of(jd, "keyword"),
                "tbUrl": tb.target if tb else None,
                "jdUrl": jd.target if jd else None,
            }
        )

    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False, indent=1)

    total = sum((i["refPrice"] or 0) * (i["qty"] if isinstance(i["qty"], int) else 1) for i in items)
    print(f"items={len(items)}  ref total={total:.2f}")
    miss = [i["id"] for i in items if not i["tbQuery"] and not i["jdQuery"]]
    if miss:
        print("no search query:", miss)


if __name__ == "__main__":
    main()
